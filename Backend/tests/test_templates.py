"""
Tests for PDF templates, custom fields, template-based documents and exports.

The LLM gateway is mocked here; the real integration is exercised manually
(see README). Every test runs against a temporary SQLite database.
"""
import gc
import io
import json
import os
import sys
import tempfile
import unittest
from pathlib import Path
from unittest import mock

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

import pymupdf  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402

from app import create_app  # noqa: E402
from config import Config  # noqa: E402
from services import template_extraction_service as extraction  # noqa: E402
from services.document_service import document_service  # noqa: E402
from services.export_service import export_filename, html_to_blocks  # noqa: E402
from services.template_service import TemplateService, template_service  # noqa: E402

LETTER_TEXT = (
    "From: Giri Nadhan\nABC Technologies, Chennai\n\nTo: The Manager\nABC Technologies\n\n"
    "Subject: Request for Leave\n\nRespected Sir/Madam,\n\n"
    "I request leave for two days due to personal reasons.\n"
)

# What a careless model might return: values attached to labels, aliases, fences
LLM_LETTER_RESPONSE = """```json
{"template_name": "Leave Letter", "document_type": "Letter", "fields": [
  {"heading": "From: Giri Nadhan", "field_type": "field", "order": 1, "level": 1, "parent_order": null},
  {"heading": "To:", "field_type": "label", "order": 2, "level": 1, "parent_order": null},
  {"heading": "Subject: Request for Leave", "field_type": "field", "order": 3, "level": 1, "parent_order": null},
  {"heading": "Respected Sir/Madam", "field_type": "salutation", "order": 4, "level": 1, "parent_order": null},
  {"heading": "Body", "field_type": "paragraph", "order": 5, "level": 1, "parent_order": null}
]}
```"""


def make_text_pdf(text: str) -> bytes:
    doc = pymupdf.open()
    page = doc.new_page()
    page.insert_textbox(pymupdf.Rect(50, 50, 550, 800), text, fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_scanned_pdf(text: str) -> bytes:
    src = pymupdf.open(stream=make_text_pdf(text), filetype='pdf')
    out = pymupdf.open()
    for page in src:
        pix = page.get_pixmap(dpi=72)
        new = out.new_page(width=page.rect.width, height=page.rect.height)
        new.insert_image(new.rect, stream=pix.tobytes('png'))
    data = out.tobytes()
    out.close()
    src.close()
    return data


def mock_client(content: str = LLM_LETTER_RESPONSE, error: dict = None):
    client = mock.MagicMock()
    client.is_configured = True
    response = error if error is not None else {'content': content, 'usage': {}}
    client.query_llm.return_value = response
    client.query_with_image.return_value = response
    return client


class TempDbTestCase(unittest.TestCase):
    def setUp(self):
        fd, self.db_path = tempfile.mkstemp(suffix='.db')
        os.close(fd)
        self._orig = (document_service.db_path, template_service.db_path)
        document_service.db_path = self.db_path
        template_service.db_path = self.db_path
        document_service._init_db()
        template_service._init_db()
        self.app = create_app()
        self.app.config['TESTING'] = True
        self.client = self.app.test_client()

    def tearDown(self):
        document_service.db_path, template_service.db_path = self._orig
        gc.collect()  # release sqlite handles before deleting on Windows
        try:
            os.remove(self.db_path)
        except OSError:
            pass

    def upload(self, pdf_bytes: bytes, llm=None, filename='leave.pdf', **form):
        llm = llm or mock_client()
        with mock.patch.object(extraction, 'get_llm_client', return_value=llm):
            data = {'file': (io.BytesIO(pdf_bytes), filename), **form}
            return self.client.post('/api/templates/upload', data=data, content_type='multipart/form-data'), llm

    def create_template(self):
        resp, _ = self.upload(make_text_pdf(LETTER_TEXT))
        self.assertEqual(resp.status_code, 201, resp.get_json())
        return resp.get_json()['template']


class NormalizationTests(unittest.TestCase):
    def test_strips_values_and_fences(self):
        result = extraction.normalize_extraction(extraction.parse_llm_json(LLM_LETTER_RESPONSE))
        headings = [f['heading'] for f in result['fields']]
        self.assertEqual(headings, ['From', 'To', 'Subject', 'Respected Sir/Madam', 'Body'])
        self.assertEqual([f['field_type'] for f in result['fields']],
                         ['field', 'field', 'field', 'section', 'paragraph'])
        self.assertNotIn('Giri', json.dumps(result))

    def test_hierarchy_from_parent_order_and_level(self):
        data = {'template_name': 'Legal Opinion', 'fields': [
            {'heading': 'Devolution of Properties', 'field_type': 'section', 'order': 1, 'level': 1},
            {'heading': 'Prior Title', 'field_type': 'section', 'order': 2, 'level': 2, 'parent_order': 1},
            {'heading': 'Present Title', 'field_type': 'section', 'order': 3, 'level': 2},  # parent inferred
            {'heading': 'Opinion', 'field_type': 'section', 'order': 4, 'level': 1},
        ]}
        fields = extraction.normalize_extraction(data)['fields']
        self.assertEqual([f['parent_index'] for f in fields], [None, 0, 0, None])

    def test_order_respected_and_sentences_dropped(self):
        data = {'fields': [
            {'heading': 'Opinion', 'order': 2},
            {'heading': 'Description of the Properties', 'order': 1},
            {'heading': ' '.join(['word'] * 40), 'order': 3},
        ]}
        fields = extraction.normalize_extraction(data)['fields']
        self.assertEqual([f['heading'] for f in fields], ['Description of the Properties', 'Opinion'])

    def test_label_value_split_only_for_short_labels(self):
        self.assertEqual(extraction.clean_extracted_heading('Schedule A: Description of the Property Rights'),
                         'Schedule A')  # short label -> treated as label: value
        self.assertEqual(extraction.clean_extracted_heading('Section 4 of the agreement between the parties: scope'),
                         'Section 4 of the agreement between the parties: scope')

    def test_invalid_json_and_empty(self):
        with self.assertRaises(extraction.ExtractionError):
            extraction.parse_llm_json('not json at all')
        with self.assertRaises(extraction.ExtractionError):
            extraction.normalize_extraction({'fields': []})
        with self.assertRaises(extraction.ExtractionError):
            extraction.normalize_extraction({'fields': 'nope'})


class UploadTests(TempDbTestCase):
    def test_upload_extracts_structure_only(self):
        resp, llm = self.upload(make_text_pdf(LETTER_TEXT))
        self.assertEqual(resp.status_code, 201)
        body = resp.get_json()
        template = body['template']
        self.assertEqual(template['template_name'], 'Leave Letter')
        self.assertEqual(body['extraction']['method'], 'text')
        self.assertEqual([f['heading'] for f in template['fields']],
                         ['From', 'To', 'Subject', 'Respected Sir/Madam', 'Body'])
        self.assertTrue(all(not f['is_custom'] for f in template['fields']))
        # The PDF text was sent to the LLM, but none of it is stored
        sent = json.dumps(llm.query_llm.call_args.args[2])
        self.assertIn('Giri Nadhan', sent)
        with open(self.db_path, 'rb') as fh:
            self.assertNotIn(b'Giri', fh.read())

    def test_scanned_pdf_uses_vision(self):
        resp, llm = self.upload(make_scanned_pdf(LETTER_TEXT))
        self.assertEqual(resp.status_code, 201, resp.get_json())
        self.assertEqual(resp.get_json()['extraction']['method'], 'vision')
        llm.query_with_image.assert_called_once()
        image_paths = llm.query_with_image.call_args.args[3]
        self.assertEqual(len(image_paths), 1)
        self.assertFalse(any(os.path.exists(p) for p in image_paths), 'temporary page images must be removed')

    def test_template_name_override(self):
        resp, _ = self.upload(make_text_pdf(LETTER_TEXT), template_name='My Leave Form')
        self.assertEqual(resp.get_json()['template']['template_name'], 'My Leave Form')

    def test_rejects_non_pdf(self):
        resp, llm = self.upload(b'plain text', filename='notes.txt')
        self.assertEqual(resp.status_code, 400)
        resp, llm = self.upload(b'plain text pretending', filename='fake.pdf')
        self.assertEqual(resp.status_code, 422)
        self.assertIn('not a valid PDF', resp.get_json()['error'])
        llm.query_llm.assert_not_called()

    def test_missing_file(self):
        resp = self.client.post('/api/templates/upload', data={}, content_type='multipart/form-data')
        self.assertEqual(resp.status_code, 400)

    def test_size_limit(self):
        with mock.patch.object(Config, 'TEMPLATE_PDF_MAX_BYTES', 200):
            resp, llm = self.upload(make_text_pdf(LETTER_TEXT))
        self.assertEqual(resp.status_code, 413)
        llm.query_llm.assert_not_called()

    def test_llm_gateway_error_is_reported(self):
        resp, _ = self.upload(make_text_pdf(LETTER_TEXT), llm=mock_client(error={'error': 'quota exceeded', 'status_code': 429}))
        self.assertEqual(resp.status_code, 502)
        self.assertIn('quota exceeded', resp.get_json()['error'])
        self.assertEqual(self.client.get('/api/templates').get_json()['count'], 0)

    def test_llm_invalid_json(self):
        resp, _ = self.upload(make_text_pdf(LETTER_TEXT), llm=mock_client(content='Sorry, I cannot help.'))
        self.assertEqual(resp.status_code, 502)
        self.assertIn('valid JSON', resp.get_json()['error'])

    def test_llm_not_configured(self):
        llm = mock_client()
        llm.is_configured = False
        resp, _ = self.upload(make_text_pdf(LETTER_TEXT), llm=llm)
        self.assertEqual(resp.status_code, 503)
        llm.query_llm.assert_not_called()


class TemplateCrudTests(TempDbTestCase):
    def test_list_get_rename_delete(self):
        template = self.create_template()
        tid = template['id']
        listing = self.client.get('/api/templates').get_json()
        self.assertEqual(listing['count'], 1)
        self.assertEqual(listing['templates'][0]['field_count'], 5)

        resp = self.client.put(f'/api/templates/{tid}', json={'template_name': '  Leave   Request '})
        self.assertEqual(resp.get_json()['template']['template_name'], 'Leave Request')
        self.assertEqual(self.client.put(f'/api/templates/{tid}', json={'template_name': ''}).status_code, 400)

        self.assertEqual(self.client.get('/api/templates/missing').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/templates/{tid}').status_code, 200)
        self.assertEqual(self.client.get(f'/api/templates/{tid}').status_code, 404)
        self.assertEqual(self.client.delete(f'/api/templates/{tid}').status_code, 404)

    def test_persists_across_restart(self):
        tid = self.create_template()['id']
        fresh = TemplateService(self.db_path)  # simulates a backend restart
        reloaded = fresh.get_template(tid)
        self.assertEqual([f['heading'] for f in reloaded['fields']],
                         ['From', 'To', 'Subject', 'Respected Sir/Madam', 'Body'])


class CustomFieldTests(TempDbTestCase):
    def test_add_edit_reorder_remove(self):
        template = self.create_template()
        tid = template['id']
        subject = next(f for f in template['fields'] if f['heading'] == 'Subject')

        # Append at end
        r = self.client.post(f'/api/templates/{tid}/fields', json={'heading': 'Employee ID'})
        self.assertEqual(r.status_code, 201)
        emp = r.get_json()['field']
        self.assertTrue(emp['is_custom'])
        self.assertEqual(emp['field_order'], 6)

        # Insert after Subject
        r = self.client.post(f'/api/templates/{tid}/fields',
                             json={'heading': 'Leave Start Date', 'field_type': 'field', 'after_field_id': subject['id']})
        headings = [f['heading'] for f in r.get_json()['template']['fields']]
        self.assertEqual(headings, ['From', 'To', 'Subject', 'Leave Start Date', 'Respected Sir/Madam', 'Body', 'Employee ID'])

        # Explicit position
        r = self.client.post(f'/api/templates/{tid}/fields', json={'heading': 'Department', 'position': 1})
        dept = r.get_json()['field']
        self.assertEqual(r.get_json()['template']['fields'][0]['heading'], 'Department')

        # Rename + move to the end
        r = self.client.put(f'/api/templates/{tid}/fields/{dept["id"]}', json={'heading': 'Dept.', 'position': 99})
        fields = r.get_json()['template']['fields']
        self.assertEqual(fields[-1]['heading'], 'Dept.')
        self.assertEqual([f['field_order'] for f in fields], list(range(1, len(fields) + 1)))

        # Validation
        self.assertEqual(self.client.post(f'/api/templates/{tid}/fields', json={'heading': '  '}).status_code, 400)
        self.assertEqual(self.client.post(f'/api/templates/{tid}/fields',
                                          json={'heading': 'X', 'field_type': 'bogus'}).status_code, 400)
        self.assertEqual(self.client.put(f'/api/templates/{tid}/fields/nope', json={'heading': 'X'}).status_code, 404)

        # Remove
        r = self.client.delete(f'/api/templates/{tid}/fields/{emp["id"]}')
        self.assertEqual(r.status_code, 200)
        self.assertNotIn('Employee ID', [f['heading'] for f in r.get_json()['template']['fields']])

        # Persisted
        reloaded = TemplateService(self.db_path).get_template(tid)
        self.assertEqual([f['heading'] for f in reloaded['fields']],
                         ['From', 'To', 'Subject', 'Leave Start Date', 'Respected Sir/Madam', 'Body', 'Dept.'])
        self.assertEqual(reloaded['custom_field_count'], 2)

    def test_nested_custom_field(self):
        template = self.create_template()
        tid = template['id']
        salutation = next(f for f in template['fields'] if f['heading'] == 'Respected Sir/Madam')
        r = self.client.post(f'/api/templates/{tid}/fields',
                             json={'heading': 'Reason', 'parent_field_id': salutation['id']})
        field = r.get_json()['field']
        self.assertEqual(field['parent_field_id'], salutation['id'])
        self.assertEqual(field['level'], 2)
        headings = [f['heading'] for f in r.get_json()['template']['fields']]
        self.assertEqual(headings.index('Reason'), headings.index('Respected Sir/Madam') + 1)
        # Cycles are rejected
        r = self.client.put(f'/api/templates/{tid}/fields/{salutation["id"]}', json={'parent_field_id': field['id']})
        self.assertEqual(r.status_code, 400)

    def test_fields_are_isolated_per_template(self):
        first = self.create_template()['id']
        second = self.create_template()['id']
        self.client.post(f'/api/templates/{first}/fields', json={'heading': 'Employee ID'})
        other = self.client.get(f'/api/templates/{second}').get_json()['template']
        self.assertNotIn('Employee ID', [f['heading'] for f in other['fields']])
        # A field id from one template can't be edited through another
        field_id = self.client.get(f'/api/templates/{first}').get_json()['template']['fields'][-1]['id']
        self.assertEqual(self.client.delete(f'/api/templates/{second}/fields/{field_id}').status_code, 404)


FILLED_DOCUMENT = """
<div data-template-field="" data-field-id="f1" data-heading="From" data-field-type="field" data-level="1" data-custom="false">
  <div data-field-label="" contenteditable="false">From:</div>
  <div data-field-body=""><p>Giri Nadhan</p><p>ABC Technologies, Chennai</p></div>
</div>
<div data-template-field="" data-field-id="f2" data-heading="Subject" data-field-type="field" data-level="1" data-custom="false">
  <div data-field-label="" contenteditable="false">Subject:</div>
  <div data-field-body=""><p>Request for <strong>Leave</strong></p></div>
</div>
<div data-template-field="" data-field-id="f3" data-heading="Respected Sir/Madam" data-field-type="section" data-level="1" data-custom="false">
  <div data-field-label="" contenteditable="false">Respected Sir/Madam</div>
  <div data-field-body=""><p>I request leave for two days.</p><ul><li><p>15 October</p></li><li><p>16 October</p></li></ul></div>
</div>
<div data-template-field="" data-field-id="f4" data-heading="Body" data-field-type="paragraph" data-level="1" data-custom="false">
  <div data-field-label="" contenteditable="false">Body</div>
  <div data-field-body=""><p>Thank you.</p></div>
</div>
<div data-template-field="" data-field-id="f5" data-heading="Employee ID" data-field-type="field" data-level="1" data-custom="true">
  <div data-field-label="" contenteditable="false">Employee ID:</div>
  <div data-field-body=""><p>ABC-10293</p></div>
</div>
<div data-template-field="" data-field-id="f6" data-heading="Department" data-field-type="field" data-level="1" data-custom="true">
  <div data-field-label="" contenteditable="false">Department:</div>
  <div data-field-body=""><p></p></div>
</div>
<p></p>
"""


class DocumentAndExportTests(TempDbTestCase):
    def test_document_from_template_lifecycle_and_export(self):
        tid = self.create_template()['id']
        r = self.client.post(f'/api/templates/{tid}/documents', json={
            'title': 'Leave Letter - Giri.docx', 'content': FILLED_DOCUMENT, 'text_content': 'x', 'word_count': 20,
        })
        self.assertEqual(r.status_code, 201)
        doc = r.get_json()['document']
        self.assertEqual(doc['template_id'], tid)
        self.assertEqual(self.client.post('/api/templates/missing/documents', json={}).status_code, 404)

        # Update in place (no duplicates)
        self.client.put(f'/api/documents/{doc["id"]}', json={'content': FILLED_DOCUMENT.replace('Thank you.', 'Thanks!')})
        docs = self.client.get('/api/documents').get_json()['documents']
        self.assertEqual(len(docs), 1)
        self.assertEqual(docs[0]['template_id'], tid)
        self.assertEqual(self.client.get(f'/api/templates/{tid}').get_json()['template']['document_count'], 1)

        # DOCX
        r = self.client.get(f'/api/documents/{doc["id"]}/export?format=docx')
        self.assertEqual(r.status_code, 200)
        self.assertIn('Leave Letter - Giri.docx', r.headers['Content-Disposition'])
        docx_text = [p.text for p in DocxDocument(io.BytesIO(r.data)).paragraphs]
        self.assertEqual(docx_text[:3], ['From: Giri Nadhan', 'ABC Technologies, Chennai', 'Subject: Request for Leave'])
        self.assertIn('Respected Sir/Madam', docx_text)
        self.assertIn('15 October', docx_text)
        self.assertIn('Thanks!', docx_text)
        self.assertIn('Employee ID: ABC-10293', docx_text)
        self.assertIn('Department:', docx_text)
        self.assertEqual(docx_text[-1], 'Department:')  # no stray empty paragraph at the end
        self.assertNotIn('Body', docx_text)  # structural label of a free-text region is not printed

        # PDF
        r = self.client.get(f'/api/documents/{doc["id"]}/export?format=pdf')
        self.assertEqual(r.status_code, 200)
        self.assertEqual(r.mimetype, 'application/pdf')
        pdf = pymupdf.open(stream=r.data, filetype='pdf')
        pdf_text = ''.join(page.get_text() for page in pdf)
        pdf.close()
        for expected in ('From: Giri Nadhan', 'Subject: Request for Leave', 'Respected Sir/Madam',
                         'Thanks!', 'Employee ID: ABC-10293', '16 October'):
            self.assertIn(expected, pdf_text)

        # Bad requests
        self.assertEqual(self.client.get(f'/api/documents/{doc["id"]}/export?format=rtf').status_code, 400)
        self.assertEqual(self.client.get('/api/documents/missing/export?format=pdf').status_code, 404)

        # Deleting the template keeps the document, unlinking it
        r = self.client.delete(f'/api/templates/{tid}')
        self.assertEqual(r.get_json()['detached_documents'], 1)
        kept = self.client.get(f'/api/documents/{doc["id"]}').get_json()['document']
        self.assertIsNone(kept['template_id'])
        self.assertIn('Giri Nadhan', kept['content'])

    def test_long_document_paginates_without_loss(self):
        paragraphs = ''.join(f'<p>Paragraph number {i} of a long document body.</p>' for i in range(400))
        r = self.client.post('/api/documents', json={'title': 'Long', 'content': f'<h1>Report</h1>{paragraphs}'})
        doc_id = r.get_json()['document']['id']
        r = self.client.get(f'/api/documents/{doc_id}/export?format=pdf')
        pdf = pymupdf.open(stream=r.data, filetype='pdf')
        self.assertGreater(pdf.page_count, 5)
        text = ''.join(page.get_text() for page in pdf)
        pdf.close()
        for i in (0, 199, 399):
            self.assertIn(f'Paragraph number {i} ', text)

    def test_plain_document_export_still_works(self):
        r = self.client.post('/api/documents', json={
            'title': 'Notes', 'content': '<h2>Agenda</h2><ol><li><p>One</p></li></ol><table><tr><th>A</th></tr><tr><td>1</td></tr></table>'})
        doc_id = r.get_json()['document']['id']
        r = self.client.get(f'/api/documents/{doc_id}/export?format=docx')
        document = DocxDocument(io.BytesIO(r.data))
        self.assertIn('Agenda', [p.text for p in document.paragraphs])
        self.assertEqual(document.tables[0].cell(1, 0).text, '1')

    def test_export_helpers(self):
        self.assertEqual(export_filename('Document1.docx', 'pdf'), 'Document1.pdf')
        self.assertEqual(export_filename('a/b:c', 'docx'), 'a_b_c.docx')
        self.assertEqual(export_filename('', 'pdf'), 'Document.pdf')
        kinds = [b.kind for b in html_to_blocks(FILLED_DOCUMENT)]
        self.assertEqual(kinds[0], 'label')


if __name__ == '__main__':
    unittest.main()
