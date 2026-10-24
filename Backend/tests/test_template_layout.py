"""
Page layout of templates: pages, heading placement (alignment / page / position)
and pictures, for PDF and DOCX uploads, and their export to DOCX / PDF.
"""
import io
import json
import sys
import unittest
from pathlib import Path

backend_dir = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(backend_dir))

import pymupdf  # noqa: E402
from docx import Document as DocxDocument  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH  # noqa: E402
from docx.shared import Pt  # noqa: E402

from services.export_service import export_document  # noqa: E402
from tests.test_templates import TempDbTestCase, mock_client  # noqa: E402


def make_png(width=120, height=48, rgb=(0, 128, 0)) -> bytes:
    pix = pymupdf.Pixmap(pymupdf.csRGB, pymupdf.IRect(0, 0, width, height), False)
    pix.set_rect(pix.irect, rgb)
    return pix.tobytes('png')


LOGO = make_png()


def llm_fields(*fields, images=None) -> str:
    """LLM answer: everything on page 1, left-aligned - the measured layout must correct it."""
    return json.dumps({
        'template_name': 'Report', 'document_type': 'Report',
        'fields': [{'heading': h, 'field_type': t, 'order': i + 1, 'level': 1, 'parent_order': None,
                    'table': None, 'align': 'left', 'page': 1, 'y': 0.1} for i, (h, t) in enumerate(fields)],
        'images': images or [],
    })


def make_layout_pdf() -> bytes:
    doc = pymupdf.open()
    for number in range(3):
        page = doc.new_page(width=595, height=842)
        if number < 2:  # logo in the top-left corner of the first two pages
            page.insert_image(pymupdf.Rect(40, 30, 140, 70), stream=LOGO)
        if number == 0:
            title = 'REPORT TITLE'
            width = pymupdf.get_text_length(title, fontsize=16)
            page.insert_text(((595 - width) / 2, 200), title, fontsize=16)
            date = 'Date: 10.07.2026'
            page.insert_text((540 - pymupdf.get_text_length(date, fontsize=11), 250), date, fontsize=11)
            page.insert_text((60, 300), 'Introduction', fontsize=11)
            page.insert_text((60, 330), 'Some body text that is long enough to be a real line of prose here.',
                             fontsize=11)
        elif number == 1:
            page.insert_text((60, 150), 'Opinion', fontsize=11)
            page.insert_text((60, 180), 'Body text of the opinion section, also long enough to measure.',
                             fontsize=11)
            page.insert_text((540 - pymupdf.get_text_length('Advocate', fontsize=11), 400), 'Advocate',
                             fontsize=11)
        else:  # continuation page: plain text only, no heading or picture
            page.insert_text((60, 100), 'Continuation text of the body without any heading at all here.',
                             fontsize=11)
    data = doc.tobytes()
    doc.close()
    return data


def make_layout_docx() -> bytes:
    doc = DocxDocument()
    doc.sections[0].header.paragraphs[0].add_run().add_picture(io.BytesIO(LOGO), width=Pt(60))
    p = doc.add_paragraph('Annual Report')
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p = doc.add_paragraph('Date: 1 January 2026')
    p.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run().add_picture(io.BytesIO(make_png(300, 60, (200, 0, 0))), width=Pt(300))
    doc.add_paragraph('Opening text of the report.')
    doc.add_page_break()
    doc.add_paragraph('Findings')
    doc.add_paragraph('Text of the findings.')
    stream = io.BytesIO()
    doc.save(stream)
    return stream.getvalue()


class PdfLayoutTests(TempDbTestCase):
    def test_pages_alignment_and_pictures_are_measured(self):
        llm = mock_client(llm_fields(('Report Title', 'section'), ('Date', 'field'), ('Introduction', 'section'),
                                     ('Opinion', 'section'), ('Advocate', 'section')))
        resp, _ = self.upload(make_layout_pdf(), llm=llm, filename='report.pdf')
        self.assertEqual(resp.status_code, 201, resp.get_json())
        template = resp.get_json()['template']

        # Page 3 has no heading or picture, so the template has 2 pages
        self.assertEqual(template['page_count'], 2)
        self.assertAlmostEqual(template['page_width'], 595, delta=1)
        fields = template['fields']
        self.assertEqual([f['page'] for f in fields], [1, 1, 1, 2, 2])
        self.assertEqual([f['align'] for f in fields], ['center', 'right', None, None, 'right'])
        ys = [f['pos_y'] for f in fields[:3]]
        self.assertEqual(ys, sorted(ys))

        images = template['images']
        self.assertEqual([i['page'] for i in images], [1, 2])
        for image in images:
            self.assertAlmostEqual(image['x'], 40 / 595, delta=0.01)
            self.assertAlmostEqual(image['y'], 30 / 842, delta=0.01)
            self.assertAlmostEqual(image['w'], 100 / 595, delta=0.01)

        got = self.client.get(f"/api/templates/{template['id']}/images/{images[0]['id']}")
        self.assertEqual(got.status_code, 200)
        self.assertIn(got.mimetype, ('image/png', 'image/jpeg'))
        self.assertTrue(len(got.data) > 100)

    def test_scanned_pdf_uses_model_positions_and_crops_located_pictures(self):
        from tests.test_templates import make_scanned_pdf
        answer = json.loads(llm_fields(('Report Title', 'section'), ('Opinion', 'section')))
        answer['fields'][0].update(page=1, y=0.2, align='center')
        answer['fields'][1].update(page=1, y=0.5)
        answer['images'] = [{'page': 1, 'x': 0.05, 'y': 0.03, 'w': 0.2, 'h': 0.06},
                            {'page': 9, 'x': 0.1, 'y': 0.1, 'w': 0.2, 'h': 0.2}]  # no such page: ignored
        resp, llm = self.upload(make_scanned_pdf('REPORT TITLE\n\nOpinion'), llm=mock_client(json.dumps(answer)),
                                filename='scan.pdf')
        self.assertEqual(resp.status_code, 201, resp.get_json())
        llm.query_with_image.assert_called()
        template = resp.get_json()['template']
        self.assertEqual([f['page'] for f in template['fields']], [1, 1])
        self.assertEqual([f['pos_y'] for f in template['fields']], [0.2, 0.5])
        self.assertEqual([f['align'] for f in template['fields']], ['center', None])
        # The full-page scan is not a picture; only the located logo is
        self.assertEqual(len(template['images']), 1)
        self.assertAlmostEqual(template['images'][0]['x'], 0.05, delta=0.001)

    def test_custom_field_joins_the_page_of_the_field_before_it(self):
        llm = mock_client(llm_fields(('Report Title', 'section'), ('Opinion', 'section')))
        resp, _ = self.upload(make_layout_pdf(), llm=llm, filename='report.pdf')
        template = resp.get_json()['template']
        opinion = template['fields'][1]
        added = self.client.post(f"/api/templates/{template['id']}/fields",
                                 json={'heading': 'Remarks', 'after_field_id': opinion['id']}).get_json()
        self.assertEqual(added['field']['page'], opinion['page'])

    def test_template_without_text_keeps_working_without_layout(self):
        # A failing layout read must not break the upload
        from services import template_extraction_service as extraction
        llm = mock_client(llm_fields(('Report Title', 'section')))
        with unittest.mock.patch.object(extraction, 'pdf_layout', side_effect=RuntimeError('boom')):
            resp, _ = self.upload(make_layout_pdf(), llm=llm, filename='report.pdf')
        self.assertEqual(resp.status_code, 201, resp.get_json())
        template = resp.get_json()['template']
        self.assertIsNone(template['page_count'])
        self.assertEqual(template['images'], [])


class DocxLayoutTests(TempDbTestCase):
    def test_docx_alignment_pages_and_pictures(self):
        llm = mock_client(llm_fields(('Annual Report', 'section'), ('Date', 'field'), ('Findings', 'section')))
        resp, _ = self.upload(make_layout_docx(), llm=llm, filename='report.docx')
        self.assertEqual(resp.status_code, 201, resp.get_json())
        template = resp.get_json()['template']

        sent = json.dumps(llm.query_llm.call_args.args[2])
        self.assertIn('[center] Annual Report', sent)
        self.assertIn('--- Page 2 ---', sent)

        fields = template['fields']
        self.assertEqual([f['align'] for f in fields], ['center', 'right', None])
        self.assertEqual([f['page'] for f in fields], [1, 1, 2])
        self.assertEqual(template['page_count'], 2)

        images = template['images']
        # Header logo on both pages + the centred banner on page 1
        self.assertEqual(sorted(i['page'] for i in images), [1, 1, 2])
        banner = max(images, key=lambda i: i['w'])
        self.assertAlmostEqual(banner['x'] + banner['w'] / 2, 0.5, delta=0.02)


class PlacementMatchingTests(unittest.TestCase):
    def test_title_beats_running_header_that_starts_with_the_same_words(self):
        from services.template_layout_service import LayoutLine, TemplateLayout, _norm, place_fields
        lines = [('EMBEDDING BENCHMARKING REPORT JD -> Resume Retrieval', 0.09, 'left'),
                 ('Embedding Benchmarking Report', 0.34, 'center'),
                 ('Date: 10.07.2026', 0.40, 'right')]
        layout = TemplateLayout(page_count=1, page_width=612, page_height=792,
                                lines=[LayoutLine(1, y, t, a, _norm(t)) for t, y, a in lines])
        fields = [{'heading': 'Embedding Benchmarking Report', 'table_key': None},
                  {'heading': 'Date', 'table_key': None}]
        place_fields(fields, layout)
        self.assertEqual([(f['pos_y'], f['align']) for f in fields], [(0.34, 'center'), (0.40, 'right')])


class LayoutExportTests(unittest.TestCase):
    HTML = (
        '<div data-template-page data-page="1" data-width="595" data-height="842">'
        '<img data-template-image data-image-id="logo" data-x="0.07" data-w="0.2" data-h="0.05" data-gap="0.01">'
        '<div data-template-field data-heading="Title" data-field-type="section" data-align="center">'
        '<div data-field-body><p>Hello</p></div></div></div>'
        '<div data-template-page data-page="2" data-width="595" data-height="842">'
        '<img data-template-image data-image-id="logo" data-x="0.5" data-w="0.2" data-h="0.05" data-gap="0">'
        '<div data-template-field data-heading="Opinion" data-field-type="section">'
        '<div data-field-body><p>Fine</p></div></div></div>'
    )

    @staticmethod
    def loader(image_id):
        return {'mime': 'image/png', 'data': LOGO} if image_id == 'logo' else None

    def test_pdf_has_one_page_per_template_page_with_pictures_in_place(self):
        pdf = pymupdf.open(stream=export_document(self.HTML, 't', 'pdf', image_loader=self.loader), filetype='pdf')
        self.assertEqual(pdf.page_count, 2)
        self.assertAlmostEqual(pdf[0].rect.width, 595, delta=1)
        first = pdf[0].get_image_info()
        second = pdf[1].get_image_info()
        self.assertEqual((len(first), len(second)), (1, 1))
        self.assertAlmostEqual(first[0]['bbox'][0], 0.07 * 595, delta=3)
        self.assertAlmostEqual(first[0]['bbox'][2] - first[0]['bbox'][0], 0.2 * 595, delta=3)
        self.assertAlmostEqual(second[0]['bbox'][0], 0.5 * 595, delta=3)
        self.assertIn('Opinion', pdf[1].get_text())

    def test_docx_has_page_break_and_pictures(self):
        doc = DocxDocument(io.BytesIO(export_document(self.HTML, 't', 'docx', image_loader=self.loader)))
        self.assertEqual(len(doc.inline_shapes), 2)
        self.assertAlmostEqual(doc.sections[0].page_width.pt, 595, delta=1)
        breaks = [p for p in doc.paragraphs if p.paragraph_format.page_break_before]
        self.assertEqual(len(breaks), 1)

    def test_documents_without_pages_export_as_before(self):
        html = '<p>Plain</p>'
        pdf = pymupdf.open(stream=export_document(html, 't', 'pdf'), filetype='pdf')
        self.assertAlmostEqual(pdf[0].rect.width, pymupdf.paper_rect('a4').width, delta=1)


if __name__ == '__main__':
    unittest.main()
