"""
Export saved editor documents (Tiptap HTML) to DOCX and PDF.

The stored HTML is parsed once into a small block model (headings,
paragraphs, list items, quotes, tables, rules) and both writers render from
that model, so DOCX and PDF output share the same structure. Template
fields (<div data-template-field>) are rendered as their heading/label
followed by the user's content.
"""
import html
import io
import logging
import re
from dataclasses import dataclass, field
from html.parser import HTMLParser
from typing import List, Optional

import pymupdf
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_COLOR_INDEX
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)


class ExportError(Exception):
    """Export failed; message is safe to show to the user."""


# ------------------------------------------------------------------ DOM parse

VOID_TAGS = {'br', 'hr', 'img', 'input', 'meta', 'link', 'col', 'wbr'}


class Node:
    def __init__(self, tag: str, attrs: dict = None, parent=None):
        self.tag = tag
        self.attrs = attrs or {}
        self.parent = parent
        self.children: list = []  # Node | str

    def find_child(self, predicate) -> Optional['Node']:
        for child in self.children:
            if isinstance(child, Node) and predicate(child):
                return child
        return None


class _TreeBuilder(HTMLParser):
    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.root = Node('root')
        self.current = self.root

    def handle_starttag(self, tag, attrs):
        node = Node(tag, {k: (v if v is not None else '') for k, v in attrs}, self.current)
        self.current.children.append(node)
        if tag not in VOID_TAGS:
            self.current = node

    def handle_startendtag(self, tag, attrs):
        node = Node(tag, {k: (v if v is not None else '') for k, v in attrs}, self.current)
        self.current.children.append(node)

    def handle_endtag(self, tag):
        node = self.current
        while node is not self.root and node.tag != tag:
            node = node.parent
        if node is not self.root:
            self.current = node.parent

    def handle_data(self, data):
        self.current.children.append(data)


def parse_html(source: str) -> Node:
    builder = _TreeBuilder()
    builder.feed(source or '')
    builder.close()
    return builder.root


# --------------------------------------------------------------- block model

@dataclass
class Run:
    text: str
    bold: bool = False
    italic: bool = False
    underline: bool = False
    strike: bool = False
    sub: bool = False
    sup: bool = False
    code: bool = False
    color: Optional[str] = None       # '#rrggbb'
    highlight: Optional[str] = None   # '#rrggbb'
    size_pt: Optional[float] = None
    font: Optional[str] = None
    href: Optional[str] = None
    is_break: bool = False


@dataclass
class Block:
    kind: str  # heading | paragraph | label | list_item | quote | table | hr | code
    runs: List[Run] = field(default_factory=list)
    level: int = 1             # heading level, or list nesting depth
    align: Optional[str] = None
    ordered: bool = False
    number: int = 0
    line_height: Optional[str] = None
    rows: list = field(default_factory=list)  # table: list[list[list[Run]]], header flags in header_rows
    header_rows: set = field(default_factory=set)


def _parse_style(style: str) -> dict:
    result = {}
    for part in (style or '').split(';'):
        if ':' in part:
            key, value = part.split(':', 1)
            result[key.strip().lower()] = value.strip()
    return result


_NAMED_COLORS = {'black': '#000000', 'white': '#ffffff', 'red': '#ff0000', 'green': '#008000', 'blue': '#0000ff',
                 'yellow': '#ffff00', 'gray': '#808080', 'grey': '#808080', 'orange': '#ffa500', 'purple': '#800080'}


def _to_hex(color: Optional[str]) -> Optional[str]:
    if not color:
        return None
    color = color.strip().lower()
    if color in _NAMED_COLORS:
        return _NAMED_COLORS[color]
    m = re.fullmatch(r'#([0-9a-f]{3})', color)
    if m:
        return '#' + ''.join(c * 2 for c in m.group(1))
    if re.fullmatch(r'#[0-9a-f]{6}', color):
        return color
    m = re.fullmatch(r'rgba?\(\s*(\d+)\s*,\s*(\d+)\s*,\s*(\d+)(?:\s*,\s*[\d.]+)?\s*\)', color)
    if m:
        return '#' + ''.join(f'{min(int(v), 255):02x}' for v in m.groups())
    return None


def _to_pt(size: Optional[str]) -> Optional[float]:
    if not size:
        return None
    m = re.fullmatch(r'([\d.]+)\s*(pt|px)?', size.strip().lower())
    if not m:
        return None
    value = float(m.group(1))
    return round(value * 0.75, 1) if m.group(2) == 'px' else value


class _BlockBuilder:
    """Walks the parsed DOM and produces a flat list of Blocks."""

    def __init__(self):
        self.blocks: List[Block] = []

    # ---- inline
    def _inline(self, node, fmt: dict, runs: List[Run]):
        for child in node.children:
            if isinstance(child, str):
                text = child.replace('\r', '').replace('\n', ' ')
                if text:
                    runs.append(Run(text=text, **fmt))
                continue
            tag = child.tag
            if tag == 'br':
                runs.append(Run(text='', is_break=True))
                continue
            if tag in ('img',):
                alt = child.attrs.get('alt')
                if alt:
                    runs.append(Run(text=f'[{alt}]', **fmt))
                continue
            new = dict(fmt)
            if tag in ('strong', 'b'):
                new['bold'] = True
            elif tag in ('em', 'i'):
                new['italic'] = True
            elif tag == 'u':
                new['underline'] = True
            elif tag in ('s', 'strike', 'del'):
                new['strike'] = True
            elif tag == 'sub':
                new['sub'] = True
            elif tag == 'sup':
                new['sup'] = True
            elif tag == 'code':
                new['code'] = True
            elif tag == 'mark':
                style = _parse_style(child.attrs.get('style', ''))
                new['highlight'] = _to_hex(child.attrs.get('data-color') or style.get('background-color')) or '#fef08a'
            elif tag == 'a':
                new['href'] = child.attrs.get('href')
                new['underline'] = True
            style = _parse_style(child.attrs.get('style', ''))
            if style.get('color'):
                new['color'] = _to_hex(style['color'])
            if style.get('font-size'):
                new['size_pt'] = _to_pt(style['font-size']) or new.get('size_pt')
            if style.get('font-family'):
                new['font'] = style['font-family'].split(',')[0].strip().strip('"\'')
            if style.get('background-color') and tag != 'mark':
                new['highlight'] = _to_hex(style['background-color'])
            self._inline(child, new, runs)

    def _runs_of(self, node) -> List[Run]:
        runs: List[Run] = []
        self._inline(node, {}, runs)
        # Trim leading/trailing whitespace of the paragraph
        while runs and not runs[0].is_break and not runs[0].text.strip():
            runs.pop(0)
        if runs and not runs[0].is_break:
            runs[0].text = runs[0].text.lstrip()
        if runs and not runs[-1].is_break:
            runs[-1].text = runs[-1].text.rstrip()
        return runs

    @staticmethod
    def _align(node) -> Optional[str]:
        align = _parse_style(node.attrs.get('style', '')).get('text-align')
        return align if align in ('left', 'center', 'right', 'justify') else None

    # ---- blocks
    def walk(self, node, list_depth: int = 0, in_quote: bool = False, heading_offset: int = 0):
        for child in node.children:
            if isinstance(child, str):
                if child.strip():
                    self.blocks.append(Block('quote' if in_quote else 'paragraph', [Run(child.strip())]))
                continue
            tag = child.tag
            if 'data-template-table' in child.attrs:
                self._template_table(child, list_depth, in_quote)
            elif 'data-template-field' in child.attrs:
                self._template_field(child, list_depth, in_quote)
            elif tag in ('h1', 'h2', 'h3', 'h4', 'h5', 'h6'):
                level = min(int(tag[1]) + heading_offset, 6)
                self.blocks.append(Block('heading', self._runs_of(child), level=level, align=self._align(child)))
            elif tag == 'p':
                style = _parse_style(child.attrs.get('style', ''))
                self.blocks.append(Block('quote' if in_quote else 'paragraph', self._runs_of(child),
                                         align=self._align(child), line_height=style.get('line-height')))
            elif tag in ('ul', 'ol'):
                self._list(child, tag == 'ol', list_depth + 1)
            elif tag == 'blockquote':
                self.walk(child, list_depth, True, heading_offset)
            elif tag == 'pre':
                text = ''.join(self._text(child)).rstrip('\n')
                for line in text.split('\n') or ['']:
                    self.blocks.append(Block('code', [Run(line, code=True)]))
            elif tag == 'hr':
                self.blocks.append(Block('hr'))
            elif tag == 'table':
                self._table(child)
            elif tag in ('img',):
                continue
            else:  # div, section, span at block level, etc.
                self.walk(child, list_depth, in_quote, heading_offset)

    def _text(self, node):
        for child in node.children:
            if isinstance(child, str):
                yield child
            elif child.tag == 'br':
                yield '\n'
            else:
                yield from self._text(child)

    def _list(self, node, ordered: bool, depth: int):
        number = int(node.attrs.get('start') or 1) if ordered and str(node.attrs.get('start', '1')).isdigit() else 1
        for item in node.children:
            if not isinstance(item, Node) or item.tag != 'li':
                continue
            first = True
            for child in item.children:
                if isinstance(child, Node) and child.tag in ('ul', 'ol'):
                    self._list(child, child.tag == 'ol', depth + 1)
                    continue
                runs = self._runs_of(child) if isinstance(child, Node) else [Run(child.strip())] if child.strip() else []
                if not runs and not first:
                    continue
                self.blocks.append(Block('list_item', runs, level=depth, ordered=ordered,
                                         number=number if first else 0))
                first = False
            if first:  # empty <li>
                self.blocks.append(Block('list_item', [], level=depth, ordered=ordered, number=number))
            number += 1

    def _table(self, node):
        rows, header_rows = [], set()
        for row in self._iter_tags(node, 'tr'):
            cells = []
            is_header = True
            for cell in row.children:
                if isinstance(cell, Node) and cell.tag in ('td', 'th'):
                    if cell.tag == 'td':
                        is_header = False
                    runs: List[Run] = []
                    for i, part in enumerate(c for c in cell.children if isinstance(c, Node)):
                        if i:
                            runs.append(Run('', is_break=True))
                        runs.extend(self._runs_of(part))
                    if not any(isinstance(c, Node) for c in cell.children):
                        runs = self._runs_of(cell)
                    cells.append(runs)
            if cells:
                if is_header:
                    header_rows.add(len(rows))
                rows.append(cells)
        if rows:
            self.blocks.append(Block('table', rows=rows, header_rows=header_rows))

    def _iter_tags(self, node, tag):
        for child in node.children:
            if isinstance(child, Node):
                if child.tag == tag:
                    yield child
                else:
                    yield from self._iter_tags(child, tag)

    def _template_table(self, node, list_depth, in_quote):
        """Column headings of a table-shaped template: header row of labels, one row of content cells."""
        header, values, extra_rows = [], [], []

        def cell_runs(container) -> List[Run]:
            start = len(self.blocks)
            self.walk(container, list_depth, in_quote)
            cell_blocks = self.blocks[start:]
            del self.blocks[start:]
            runs: List[Run] = []
            for block in cell_blocks:
                if not _has_text(block.runs):
                    continue
                if runs:
                    runs.append(Run('', is_break=True))
                runs.extend(block.runs)
            return runs

        for column in node.children:
            if isinstance(column, Node) and 'data-template-table-row' in column.attrs:
                extra_rows.append([cell_runs(cell) for cell in column.children
                                   if isinstance(cell, Node) and 'data-template-table-cell' in cell.attrs])
                continue
            if not isinstance(column, Node) or 'data-template-field' not in column.attrs:
                continue
            body = column.find_child(lambda n: 'data-field-body' in n.attrs) or column
            start = len(self.blocks)
            self.walk(body, list_depth, in_quote)
            cell_blocks = self.blocks[start:]
            del self.blocks[start:]
            runs: List[Run] = []
            for block in cell_blocks:
                if not _has_text(block.runs):
                    continue
                if runs:
                    runs.append(Run('', is_break=True))
                runs.extend(block.runs)
            header.append([Run(column.attrs.get('data-heading', '').strip().rstrip(':'), bold=True)])
            values.append(runs)
        if header:
            width = len(header)
            rows = [header, values] + [(row + [[] for _ in range(width)])[:width] for row in extra_rows]
            self.blocks.append(Block('table', rows=rows, header_rows={0}))

    def _template_field(self, node, list_depth, in_quote):
        heading = node.attrs.get('data-heading', '').strip()
        field_type = node.attrs.get('data-field-type', 'field')
        try:
            level = max(1, min(int(node.attrs.get('data-level') or 1), 6))
        except ValueError:
            level = 1
        align = node.attrs.get('data-align')
        align = align if align in ('center', 'right') else None
        body = node.find_child(lambda n: 'data-field-body' in n.attrs) or node
        start = len(self.blocks)
        self.walk(body, list_depth, in_quote)
        body_blocks = [b for b in self.blocks[start:]]
        del self.blocks[start:]
        if align and field_type != 'paragraph':
            # The field's placement applies to its text unless a paragraph was aligned explicitly
            for b in body_blocks:
                if b.kind == 'paragraph' and not b.align:
                    b.align = align
        # Ignore empty paragraphs that only exist as editable placeholders
        while body_blocks and body_blocks[-1].kind == 'paragraph' and not _has_text(body_blocks[-1].runs):
            body_blocks.pop()

        if field_type == 'paragraph':
            # Free-text region ("Body"): its label is structural, not document text
            self.blocks.extend(body_blocks)
            return
        if field_type == 'field':
            label = Run(f'{heading}:' if heading and not heading.endswith(':') else heading, bold=True)
            if body_blocks and body_blocks[0].kind == 'paragraph':
                first = body_blocks.pop(0)
                runs = [label, Run(' ')] + first.runs if _has_text(first.runs) else [label]
                self.blocks.append(Block('label', runs, level=level, align=first.align or align))
            else:
                self.blocks.append(Block('label', [label], level=level, align=align))
            self.blocks.extend(body_blocks)
            return
        # section
        self.blocks.append(Block('section', [Run(heading, bold=True)], level=level, align=align))
        self.blocks.extend(body_blocks)


def _has_text(runs: List[Run]) -> bool:
    return any(r.text.strip() for r in runs if not r.is_break)


def html_to_blocks(source: str) -> List[Block]:
    builder = _BlockBuilder()
    builder.walk(parse_html(source))
    # The editor keeps an empty trailing paragraph after the last block; don't export it
    while builder.blocks and builder.blocks[-1].kind == 'paragraph' and not _has_text(builder.blocks[-1].runs):
        builder.blocks.pop()
    for block in builder.blocks:
        _fix_colors(block.runs)
        for row in block.rows:
            for cell in row:
                _fix_colors(cell)
    return builder.blocks


def _fix_colors(runs: List[Run]):
    """Text colours picked for the dark editor (near-white) are unreadable on white paper."""
    for run in runs:
        if run.color:
            r, g, b = (int(run.color[i:i + 2], 16) for i in (1, 3, 5))
            if (0.299 * r + 0.587 * g + 0.114 * b) > 200:
                run.color = None


# ------------------------------------------------------------------- DOCX

_HIGHLIGHT_MAP = [
    ((255, 255, 0), WD_COLOR_INDEX.YELLOW), ((0, 255, 0), WD_COLOR_INDEX.BRIGHT_GREEN),
    ((0, 255, 255), WD_COLOR_INDEX.TURQUOISE), ((255, 0, 255), WD_COLOR_INDEX.PINK),
    ((0, 0, 255), WD_COLOR_INDEX.BLUE), ((255, 0, 0), WD_COLOR_INDEX.RED),
    ((128, 128, 128), WD_COLOR_INDEX.GRAY_25), ((0, 128, 0), WD_COLOR_INDEX.GREEN),
    ((255, 165, 0), WD_COLOR_INDEX.YELLOW), ((200, 160, 255), WD_COLOR_INDEX.VIOLET),
]


def _nearest_highlight(hex_color: str):
    r, g, b = (int(hex_color[i:i + 2], 16) for i in (1, 3, 5))
    return min(_HIGHLIGHT_MAP, key=lambda item: sum((a - c) ** 2 for a, c in zip(item[0], (r, g, b))))[1]


_DOCX_ALIGN = {'left': WD_ALIGN_PARAGRAPH.LEFT, 'center': WD_ALIGN_PARAGRAPH.CENTER,
               'right': WD_ALIGN_PARAGRAPH.RIGHT, 'justify': WD_ALIGN_PARAGRAPH.JUSTIFY}


def _add_docx_runs(paragraph, runs: List[Run]):
    for run in runs:
        if run.is_break:
            paragraph.add_run().add_break()
            continue
        if not run.text:
            continue
        r = paragraph.add_run(run.text)
        r.bold = run.bold or None
        r.italic = run.italic or None
        r.underline = run.underline or None
        if run.strike:
            r.font.strike = True
        if run.sub:
            r.font.subscript = True
        if run.sup:
            r.font.superscript = True
        if run.code:
            r.font.name = 'Consolas'
        elif run.font:
            r.font.name = run.font
        if run.size_pt:
            r.font.size = Pt(run.size_pt)
        if run.color:
            r.font.color.rgb = RGBColor.from_string(run.color[1:].upper())
        if run.highlight:
            r.font.highlight_color = _nearest_highlight(run.highlight)


def _docx_hr(paragraph):
    p_pr = paragraph._p.get_or_add_pPr()
    borders = OxmlElement('w:pBdr')
    bottom = OxmlElement('w:bottom')
    for key, value in (('w:val', 'single'), ('w:sz', '6'), ('w:space', '1'), ('w:color', '999999')):
        bottom.set(qn(key), value)
    borders.append(bottom)
    p_pr.append(borders)


def blocks_to_docx(blocks: List[Block], title: str) -> bytes:
    doc = DocxDocument()
    normal = doc.styles['Normal']
    normal.font.name = 'Calibri'
    normal.font.size = Pt(11)
    doc.core_properties.title = title

    for block in blocks:
        if block.kind == 'heading':
            p = doc.add_heading(level=min(block.level, 9))
            _add_docx_runs(p, block.runs)
        elif block.kind == 'section':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.keep_with_next = True
            _add_docx_runs(p, [Run(r.text, bold=True, size_pt=13 if block.level == 1 else 12) for r in block.runs])
        elif block.kind == 'label':
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(4)
            _add_docx_runs(p, block.runs)
        elif block.kind == 'list_item':
            style = 'List Number' if block.ordered else 'List Bullet'
            if block.level >= 2:
                style += f' {min(block.level, 3)}'
            try:
                p = doc.add_paragraph(style=style)
            except KeyError:
                p = doc.add_paragraph(style='List Number' if block.ordered else 'List Bullet')
            _add_docx_runs(p, block.runs)
        elif block.kind == 'quote':
            p = doc.add_paragraph(style='Quote')
            _add_docx_runs(p, block.runs)
        elif block.kind == 'code':
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(0)
            _add_docx_runs(p, block.runs)
        elif block.kind == 'hr':
            _docx_hr(doc.add_paragraph())
            continue
        elif block.kind == 'table':
            columns = max(len(row) for row in block.rows)
            table = doc.add_table(rows=len(block.rows), cols=columns)
            table.style = 'Table Grid'
            for r_index, row in enumerate(block.rows):
                for c_index, cell_runs in enumerate(row):
                    cell = table.cell(r_index, c_index)
                    runs = cell_runs
                    if r_index in block.header_rows:
                        runs = [Run(**{**r.__dict__, 'bold': True}) for r in cell_runs]
                    _add_docx_runs(cell.paragraphs[0], runs)
            doc.add_paragraph()
            continue
        else:  # paragraph
            p = doc.add_paragraph()
            _add_docx_runs(p, block.runs)
            if block.line_height:
                try:
                    p.paragraph_format.line_spacing = float(block.line_height)
                except ValueError:
                    pass
        if block.align in _DOCX_ALIGN:
            p.alignment = _DOCX_ALIGN[block.align]

    buffer = io.BytesIO()
    doc.save(buffer)
    return buffer.getvalue()


# -------------------------------------------------------------------- PDF

PDF_CSS = """
* { font-family: sans-serif; }
body { font-size: 11pt; line-height: 1.45; color: #111111; }
h1 { font-size: 20pt; margin: 14pt 0 6pt 0; }
h2 { font-size: 16pt; margin: 12pt 0 5pt 0; }
h3 { font-size: 13.5pt; margin: 10pt 0 4pt 0; }
h4, h5, h6 { font-size: 12pt; margin: 8pt 0 3pt 0; }
p { margin: 0 0 6pt 0; }
p.section { font-weight: bold; font-size: 13pt; margin: 10pt 0 4pt 0; }
p.section2 { font-weight: bold; font-size: 12pt; margin: 8pt 0 3pt 0; }
p.label { margin: 5pt 0 4pt 0; }
p.quote { margin-left: 18pt; font-style: italic; color: #444444; }
p.code { font-family: monospace; font-size: 9.5pt; margin: 0; }
p.li { margin: 0 0 3pt 0; }
table { border-collapse: collapse; margin: 4pt 0 8pt 0; }
td, th { border: 1px solid #888888; padding: 3pt 5pt; vertical-align: top; }
th { font-weight: bold; background-color: #eeeeee; }
hr { margin: 6pt 0; }
"""


def _runs_to_html(runs: List[Run]) -> str:
    out = []
    for run in runs:
        if run.is_break:
            out.append('<br/>')
            continue
        if not run.text:
            continue
        text = html.escape(run.text)
        styles = []
        if run.color:
            styles.append(f'color:{run.color}')
        if run.highlight:
            styles.append(f'background-color:{run.highlight}')
        if run.size_pt:
            styles.append(f'font-size:{run.size_pt}pt')
        if run.font and not run.code:
            styles.append(f'font-family:{html.escape(run.font)}, sans-serif')
        if styles:
            text = f'<span style="{";".join(styles)}">{text}</span>'
        if run.code:
            text = f'<code>{text}</code>'
        if run.sub:
            text = f'<sub>{text}</sub>'
        if run.sup:
            text = f'<sup>{text}</sup>'
        if run.strike:
            text = f'<s>{text}</s>'
        if run.underline:
            text = f'<u>{text}</u>'
        if run.italic:
            text = f'<i>{text}</i>'
        if run.bold:
            text = f'<b>{text}</b>'
        out.append(text)
    return ''.join(out) or '&#160;'


def blocks_to_html(blocks: List[Block]) -> str:
    parts = []
    for block in blocks:
        align = f' style="text-align:{block.align}"' if block.align else ''
        if block.kind == 'heading':
            parts.append(f'<h{block.level}{align}>{_runs_to_html(block.runs)}</h{block.level}>')
        elif block.kind == 'section':
            css = 'section' if block.level == 1 else 'section2'
            indent = f' style="margin-left:{(block.level - 1) * 12}pt"' if block.level > 1 else ''
            if block.align:
                indent = f' style="{indent[8:-1] + ";" if indent else ""}text-align:{block.align}"'
            parts.append(f'<p class="{css}"{indent}>{_runs_to_html(block.runs)}</p>')
        elif block.kind == 'label':
            parts.append(f'<p class="label"{align}>{_runs_to_html(block.runs)}</p>')
        elif block.kind == 'list_item':
            if block.number:
                marker = f'{block.number}.' if block.ordered else '&#8226;'
            else:
                marker = '&#160;'
            indent = 14 + (block.level - 1) * 18
            parts.append(f'<p class="li" style="margin-left:{indent}pt">{marker}&#160;&#160;{_runs_to_html(block.runs)}</p>')
        elif block.kind == 'quote':
            parts.append(f'<p class="quote"{align}>{_runs_to_html(block.runs)}</p>')
        elif block.kind == 'code':
            parts.append(f'<p class="code">{_runs_to_html(block.runs)}</p>')
        elif block.kind == 'hr':
            parts.append('<hr/>')
        elif block.kind == 'table':
            rows = []
            for r_index, row in enumerate(block.rows):
                tag = 'th' if r_index in block.header_rows else 'td'
                rows.append('<tr>' + ''.join(f'<{tag}>{_runs_to_html(cell)}</{tag}>' for cell in row) + '</tr>')
            parts.append('<table>' + ''.join(rows) + '</table>')
        else:
            style = []
            if block.align:
                style.append(f'text-align:{block.align}')
            if block.line_height:
                style.append(f'line-height:{html.escape(block.line_height)}')
            style_attr = f' style="{";".join(style)}"' if style else ''
            parts.append(f'<p{style_attr}>{_runs_to_html(block.runs)}</p>')
    return '<html><body>' + '\n'.join(parts) + '</body></html>'


def blocks_to_pdf(blocks: List[Block], title: str) -> bytes:
    story = pymupdf.Story(html=blocks_to_html(blocks), user_css=PDF_CSS)
    buffer = io.BytesIO()
    writer = pymupdf.DocumentWriter(buffer)
    page_rect = pymupdf.paper_rect('a4')
    content_rect = page_rect + (56, 56, -56, -56)  # ~2 cm margins
    more, pages = 1, 0
    while more:
        device = writer.begin_page(page_rect)
        more, _ = story.place(content_rect)
        story.draw(device)
        writer.end_page()
        pages += 1
        if pages > 2000:
            raise ExportError('The document is too long to export as PDF.')
    writer.close()

    # Set document metadata
    pdf = pymupdf.open(stream=buffer.getvalue(), filetype='pdf')
    try:
        pdf.set_metadata({'title': title, 'creator': 'RAGSUITE Document Studio'})
        return pdf.tobytes(garbage=3, deflate=True)
    finally:
        pdf.close()


# ------------------------------------------------------------------ public

def export_document(content_html: str, title: str, fmt: str) -> bytes:
    blocks = html_to_blocks(content_html)
    if not blocks:
        blocks = [Block('paragraph', [])]
    if fmt == 'docx':
        return blocks_to_docx(blocks, title)
    if fmt == 'pdf':
        return blocks_to_pdf(blocks, title)
    raise ExportError(f'Unsupported export format: {fmt}')


def export_filename(title: str, fmt: str) -> str:
    base = (title or 'Document').strip()
    base = re.sub(r'\.(docx?|pdf|txt|html?)$', '', base, flags=re.IGNORECASE)
    base = re.sub(r'[\\/:*?"<>|\x00-\x1f]+', '_', base).strip(' ._') or 'Document'
    return f'{base[:120]}.{fmt}'
