"""
Page layout of an uploaded template: page size, which page each heading is on,
where it sits (vertical position, left / centre / right) and the pictures
(logos, banners, signatures...) with their exact position on the page.

All positions are fractions of the page (x, w of the page width; y, h of the
page height) so the editor and the exports can scale them to any size.

The LLM decides WHICH lines are headings; this module measures WHERE they are
from the source file itself (PDF text layer / DOCX paragraphs). The LLM's own
page / position guesses are used only where the file has no text to measure
(scanned pages).
"""
import logging
import re
from dataclasses import dataclass, field
from typing import Dict, List, Optional, Tuple

import pymupdf
from lxml import etree

logger = logging.getLogger(__name__)

MAX_IMAGES = 60
MAX_IMAGE_BYTES_TOTAL = 15 * 1024 * 1024
# Pictures covering more of the page than this are page scans / backgrounds, not logos
MAX_IMAGE_PAGE_AREA = 0.6
MIN_IMAGE_SIDE_PT = 12
IMAGE_RENDER_DPI = 150
BROWSER_IMAGE_TYPES = {'image/png', 'image/jpeg', 'image/gif'}

_W = '{http://schemas.openxmlformats.org/wordprocessingml/2006/main}'
_WP = '{http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing}'
_A = '{http://schemas.openxmlformats.org/drawingml/2006/main}'
_R = '{http://schemas.openxmlformats.org/officeDocument/2006/relationships}'
EMU_PER_PT = 12700


@dataclass
class LayoutImage:
    page: int          # 1-based
    x: float
    y: float
    w: float
    h: float
    mime: str
    data: bytes


@dataclass
class LayoutLine:
    """A line / paragraph of source text with its measured placement."""
    page: int
    y: float
    text: str
    align: Optional[str]   # 'left' | 'center' | 'right'
    norm: str = ''


@dataclass
class TemplateLayout:
    page_count: int
    page_width: float      # points
    page_height: float     # points
    lines: List[LayoutLine] = field(default_factory=list)
    images: List[LayoutImage] = field(default_factory=list)
    measured_pages: set = field(default_factory=set)  # pages whose text could be measured
    llm_parts: List[Tuple[int, str]] = field(default_factory=list)  # DOCX: (page, line) for the LLM


def _norm(text: str) -> str:
    return re.sub(r'[^0-9a-z]+', '', (text or '').lower())


def _clamp(value: float, low: float = 0.0, high: float = 1.0) -> float:
    return max(low, min(high, value))


def _encode_pixmap(pix: 'pymupdf.Pixmap') -> Tuple[str, bytes]:
    """PNG for small / flat pictures, JPEG for large photographic ones."""
    if pix.alpha:
        pix = pymupdf.Pixmap(pix, 0)
    png = pix.tobytes('png')
    if len(png) <= 300 * 1024:
        return 'image/png', png
    return 'image/jpeg', pix.tobytes('jpeg', jpg_quality=85)


def _add_image(layout: TemplateLayout, image: LayoutImage) -> None:
    used = sum(len(i.data) for i in layout.images)
    if len(layout.images) >= MAX_IMAGES or used + len(image.data) > MAX_IMAGE_BYTES_TOTAL:
        logger.info("Template image limit reached; skipping an image on page %s", image.page)
        return
    layout.images.append(image)


def _merge_rects(rects: List[pymupdf.Rect]) -> List[pymupdf.Rect]:
    """Union overlapping rectangles (pictures drawn as several tiles become one)."""
    merged: List[pymupdf.Rect] = []
    for rect in sorted(rects, key=lambda r: (r.y0, r.x0)):
        for index, other in enumerate(merged):
            if rect.intersects(other):
                merged[index] = other | rect
                break
        else:
            merged.append(pymupdf.Rect(rect))
    # A union can create new overlaps; repeat until stable
    if len(merged) < len(rects):
        return _merge_rects(merged)
    return merged


def _line_align(x0: float, x1: float, page_width: float, text_left: float, text_right: float) -> str:
    center = (x0 + x1) / 2
    if x0 > page_width * 0.5 and x1 >= text_right - page_width * 0.06:
        return 'right'
    if abs(center - page_width / 2) <= page_width * 0.06 and x0 > text_left + page_width * 0.08:
        return 'center'
    return 'left'


# --------------------------------------------------------------------- PDF

def pdf_layout(pdf_bytes: bytes, max_pages: int) -> TemplateLayout:
    doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    try:
        first = doc[0].rect
        layout = TemplateLayout(page_count=min(doc.page_count, max_pages),
                                page_width=first.width, page_height=first.height)
        for index in range(layout.page_count):
            page = doc[index]
            rect = page.rect
            width, height = rect.width or 1, rect.height or 1

            # Text lines with their position (only pages that have a text layer)
            raw_lines = []
            try:
                for block in page.get_text('dict')['blocks']:
                    for line in block.get('lines', []):
                        text = ''.join(span.get('text', '') for span in line.get('spans', [])).strip()
                        if text:
                            raw_lines.append((pymupdf.Rect(line['bbox']), text))
            except Exception:
                logger.debug("Text layout extraction failed on page %s", index + 1, exc_info=True)
            if sum(len(t) for _, t in raw_lines) >= 25:
                layout.measured_pages.add(index + 1)
                text_left = min(r.x0 for r, _ in raw_lines)
                text_right = max(r.x1 for r, _ in raw_lines)
                for bbox, text in sorted(raw_lines, key=lambda item: (round(item[0].y0), item[0].x0)):
                    layout.lines.append(LayoutLine(
                        page=index + 1, y=_clamp((bbox.y0 - rect.y0) / height), text=text,
                        align=_line_align(bbox.x0 - rect.x0, bbox.x1 - rect.x0, width,
                                          text_left - rect.x0, text_right - rect.x0),
                        norm=_norm(text),
                    ))

            # Embedded pictures, rendered as they appear on the page
            rects = []
            try:
                for info in page.get_image_info():
                    bbox = pymupdf.Rect(info['bbox']) & rect
                    if bbox.is_empty or bbox.width < MIN_IMAGE_SIDE_PT or bbox.height < MIN_IMAGE_SIDE_PT:
                        continue
                    if (bbox.width * bbox.height) / (width * height) > MAX_IMAGE_PAGE_AREA:
                        continue  # page scan or background
                    rects.append(bbox)
            except Exception:
                logger.debug("Image listing failed on page %s", index + 1, exc_info=True)
            for bbox in _merge_rects(rects):
                try:
                    pix = page.get_pixmap(clip=bbox, dpi=IMAGE_RENDER_DPI, alpha=False)
                    mime, data = _encode_pixmap(pix)
                except Exception:
                    logger.debug("Image render failed on page %s", index + 1, exc_info=True)
                    continue
                _add_image(layout, LayoutImage(
                    page=index + 1,
                    x=_clamp((bbox.x0 - rect.x0) / width), y=_clamp((bbox.y0 - rect.y0) / height),
                    w=_clamp(bbox.width / width), h=_clamp(bbox.height / height),
                    mime=mime, data=data,
                ))
        return layout
    finally:
        doc.close()


def crop_pdf_regions(pdf_bytes: bytes, layout: TemplateLayout, regions: List[Dict]) -> None:
    """Add pictures the vision model located on scanned pages (boxes as page fractions)."""
    if not regions:
        return
    doc = pymupdf.open(stream=pdf_bytes, filetype='pdf')
    try:
        for region in regions:
            page_no = region['page']
            if page_no < 1 or page_no > layout.page_count or page_no in layout.measured_pages:
                continue  # text pages already had their real images extracted
            page = doc[page_no - 1]
            rect = page.rect
            clip = pymupdf.Rect(rect.x0 + region['x'] * rect.width, rect.y0 + region['y'] * rect.height,
                                rect.x0 + (region['x'] + region['w']) * rect.width,
                                rect.y0 + (region['y'] + region['h']) * rect.height) & rect
            if clip.is_empty or clip.width < MIN_IMAGE_SIDE_PT or clip.height < MIN_IMAGE_SIDE_PT:
                continue
            try:
                mime, data = _encode_pixmap(page.get_pixmap(clip=clip, dpi=IMAGE_RENDER_DPI, alpha=False))
            except Exception:
                logger.debug("Region crop failed on page %s", page_no, exc_info=True)
                continue
            _add_image(layout, LayoutImage(page=page_no, x=region['x'], y=region['y'],
                                           w=region['w'], h=region['h'], mime=mime, data=data))
    finally:
        doc.close()


# -------------------------------------------------------------------- DOCX

def _docx_align(paragraph) -> str:
    """Effective paragraph alignment, following the style chain."""
    def jc_of(p_pr):
        if p_pr is None:
            return None
        jc = p_pr.find(f'{_W}jc')
        return jc.get(f'{_W}val') if jc is not None else None

    value = jc_of(paragraph._p.pPr)
    style = paragraph.style
    while value is None and style is not None:
        value = jc_of(style.element.find(f'{_W}pPr'))
        style = style.base_style
    if value == 'center':
        return 'center'
    if value in ('right', 'end'):
        return 'right'
    return 'left'


def _docx_picture(drawing, part) -> Optional[Tuple[str, bytes, float, float]]:
    """(mime, data, width_pt, height_pt) of a <w:drawing>, or None."""
    blip = drawing.find(f'.//{_A}blip')
    extent = drawing.find(f'.//{_WP}extent')
    if blip is None or extent is None:
        return None
    rel_id = blip.get(f'{_R}embed')
    image_part = part.related_parts.get(rel_id) if rel_id else None
    if image_part is None:
        return None
    mime, data = image_part.content_type, image_part.blob
    if mime not in BROWSER_IMAGE_TYPES:
        try:  # e.g. BMP / TIFF -> PNG; vector formats (EMF/WMF) cannot be converted and are skipped
            mime, data = _encode_pixmap(pymupdf.Pixmap(data))
        except Exception:
            return None
    try:
        width_pt = int(extent.get('cx')) / EMU_PER_PT
        height_pt = int(extent.get('cy')) / EMU_PER_PT
    except (TypeError, ValueError):
        return None
    if width_pt < MIN_IMAGE_SIDE_PT or height_pt < MIN_IMAGE_SIDE_PT:
        return None
    return mime, data, width_pt, height_pt


def _anchor_offset(anchor, axis: str, size_pt: float, page_pt: float, margin_start: float,
                   margin_end: float, paragraph_pt: Optional[float] = None) -> Optional[float]:
    """Absolute position (pt from the page edge) of an anchored picture, if it can be read.
    `paragraph_pt` is the top of the anchoring paragraph, for pictures placed relative to it."""
    pos = anchor.find(f'{_WP}position{axis}')
    if pos is None:
        return None
    relative = pos.get('relativeFrom', 'column')
    if relative in ('page', 'leftMargin', 'topMargin'):
        start, span = 0.0, page_pt
    elif relative in ('rightMargin', 'bottomMargin'):
        start, span = page_pt - margin_end, margin_end
    elif relative in ('paragraph', 'line') and paragraph_pt is not None:
        start, span = paragraph_pt, page_pt - margin_end - paragraph_pt
    else:
        start, span = margin_start, page_pt - margin_start - margin_end
    offset = pos.find(f'{_WP}posOffset')
    if offset is not None and offset.text and offset.text.strip().lstrip('-').isdigit():
        return start + int(offset.text) / EMU_PER_PT
    align = pos.find(f'{_WP}align')
    if align is not None and align.text:
        where = align.text.strip()
        if where in ('center',):
            return start + (span - size_pt) / 2
        if where in ('right', 'bottom', 'outside'):
            return start + span - size_pt
        return start
    return None


# Word records where its pages end (<w:lastRenderedPageBreak/>) only in files it saved itself;
# files written by other tools have just their explicit breaks. For those the text flow is
# estimated from the formatting: font sizes, line and paragraph spacing, table rows, pictures.

# Single-spacing line height (em) and character width relative to Calibri
_FONT_METRICS = {
    'calibri': (1.22, 1.0), 'calibri light': (1.22, 0.98), 'aptos': (1.2, 1.06), 'cambria': (1.17, 1.06),
    'arial': (1.15, 1.12), 'helvetica': (1.15, 1.12), 'times new roman': (1.15, 0.98), 'georgia': (1.14, 1.1),
    'segoe ui': (1.33, 1.1), 'verdana': (1.22, 1.3), 'tahoma': (1.21, 1.08), 'courier new': (1.13, 1.27),
}
_DEFAULT_FONT_METRICS = (1.17, 1.08)
DEFAULT_FONT_PT = 10.0       # Word's size when no style sets one
DEFAULT_TAB_PT = 36.0
DEFAULT_CELL_MARGIN_PT = 5.4


def _on(element) -> bool:
    """Value of a toggle property such as <w:b/> or <w:keepNext w:val="0"/>."""
    return element.get(f'{_W}val', 'true').lower() not in ('0', 'false', 'off')


def _twips(value, default: float = 0.0) -> float:
    try:
        return int(value) / 20
    except (TypeError, ValueError):
        return default


def _char_em(char: str) -> float:
    """Approximate Calibri advance width of a character (em)."""
    if char.isupper():
        return 0.58
    if char.isdigit():
        return 0.507
    if char.isalpha():
        return 0.46
    return 0.35


@dataclass
class _Box:
    """Estimated vertical extent of a paragraph (points)."""
    before: float
    height: float
    after: float
    line: float
    keep_next: bool = False
    border_key: str = ''        # paragraphs with the same border are drawn as one bordered box
    border_top: float = 0.0
    border_bottom: float = 0.0


def _gap(previous: Optional[_Box], box: _Box) -> float:
    """Distance from the end of `previous` (its space after included) to the top of `box`.
    Word uses the larger of the two paragraph spacings; a shared border is drawn once."""
    if previous is None:
        return box.before
    gap = max(box.before - previous.after, 0.0)
    if box.border_key and box.border_key == previous.border_key:
        gap -= previous.border_bottom + box.border_top
    return gap


def _stack(boxes: List[_Box]) -> Tuple[List[float], float]:
    """Top of each paragraph stacked one under the other, and the total height."""
    offset, previous, tops = 0.0, None, []
    for box in boxes:
        offset += _gap(previous, box)
        tops.append(offset)
        offset += box.height + box.after
        previous = box
    return tops, offset


class _DocxFlow:
    """Estimated heights of DOCX paragraphs, table rows and headers / footers."""

    def __init__(self, document):
        defaults = document.styles.element.find(f'{_W}docDefaults')
        self.default_ppr = defaults.find(f'{_W}pPrDefault/{_W}pPr') if defaults is not None else None
        self.default_rpr = defaults.find(f'{_W}rPrDefault/{_W}rPr') if defaults is not None else None
        self.theme_fonts = {'minor': 'calibri', 'major': 'calibri light'}
        try:
            from docx.opc.constants import RELATIONSHIP_TYPE as RT
            theme = etree.fromstring(document.part.part_related_by(RT.THEME).blob)
            for kind in ('minor', 'major'):
                latin = theme.find(f'.//{_A}{kind}Font/{_A}latin')
                if latin is not None and latin.get('typeface'):
                    self.theme_fonts[kind] = latin.get('typeface').lower()
        except Exception:
            logger.debug("DOCX theme fonts could not be read", exc_info=True)

    @staticmethod
    def _first(chain, child: str, attr: Optional[str] = None):
        """Nearest `child` element in a property chain (that sets `attr`, if given)."""
        for props in chain:
            element = props.find(f'{_W}{child}')
            if element is not None and (attr is None or element.get(f'{_W}{attr}') is not None):
                return element
        return None

    def _chains(self, paragraph) -> Tuple[list, list]:
        """(pPr, rPr) chains of a paragraph: its own, its style and base styles, the document defaults."""
        p_pr = [paragraph._p.pPr] if paragraph._p.pPr is not None else []
        r_pr = []
        try:
            style = paragraph.style
        except Exception:
            style = None
        while style is not None:
            for tag, chain in (('pPr', p_pr), ('rPr', r_pr)):
                element = style.element.find(f'{_W}{tag}')
                if element is not None:
                    chain.append(element)
            style = style.base_style
        p_pr += [self.default_ppr] if self.default_ppr is not None else []
        r_pr += [self.default_rpr] if self.default_rpr is not None else []
        return p_pr, r_pr

    def _font(self, chain) -> Tuple[float, float, float]:
        """(size, single line height, width factor) for a run property chain."""
        size = DEFAULT_FONT_PT
        sz = self._first(chain, 'sz', 'val')
        if sz is not None and sz.get(f'{_W}val').isdigit():
            size = int(sz.get(f'{_W}val')) / 2
        name = ''
        for props in chain:
            fonts = props.find(f'{_W}rFonts')
            if fonts is None:
                continue
            if fonts.get(f'{_W}ascii'):
                name = fonts.get(f'{_W}ascii').lower()
                break
            if fonts.get(f'{_W}asciiTheme'):
                name = self.theme_fonts['major' if fonts.get(f'{_W}asciiTheme').startswith('major') else 'minor']
                break
        line_em, width = _FONT_METRICS.get(name, _DEFAULT_FONT_METRICS)
        bold = self._first(chain, 'b')
        if bold is not None and _on(bold):
            width *= 1.05
        return size, size * line_em, width

    def paragraph_box(self, paragraph, width: float) -> _Box:
        p = paragraph._p
        p_chain, r_chain = self._chains(paragraph)

        def spacing(attr):
            element = self._first(p_chain, 'spacing', attr)
            return element.get(f'{_W}{attr}') if element is not None else None

        def indent(*attrs):
            for attr in attrs:
                element = self._first(p_chain, 'ind', attr)
                if element is not None:
                    return _twips(element.get(f'{_W}{attr}'))
            return 0.0

        avail = max(36.0, width - indent('left', 'start') - indent('right', 'end'))
        mark = p.pPr.find(f'{_W}rPr') if p.pPr is not None else None
        _, mark_line, _ = self._font(([mark] if mark is not None else []) + r_chain)

        # Greedy word wrap over the runs
        lines, x, word, text_line, picture_h = 1, indent('firstLine') - indent('hanging'), 0.0, 0.0, 0.0

        def place_word():
            nonlocal lines, x, word
            if x + word > avail and x > 0:
                lines, x = lines + 1, 0.0
            while word > avail:
                lines, word = lines + 1, word - avail
            x, word = x + word, 0.0

        for run in p.iter(f'{_W}r'):
            if next(run.iterancestors(f'{_W}txbxContent'), None) is not None:
                continue  # text box contents float outside the paragraph flow
            r_pr = run.find(f'{_W}rPr')
            size, run_line, factor = self._font(([r_pr] if r_pr is not None else []) + r_chain)
            for child in run:
                if child.tag == f'{_W}t':
                    if child.text:
                        text_line = max(text_line, run_line)
                    for char in child.text or '':
                        if char == ' ':
                            place_word()
                            x += 0.226 * size * factor
                        else:
                            word += _char_em(char) * size * factor
                elif child.tag == f'{_W}tab':
                    place_word()
                    x = (int(x / DEFAULT_TAB_PT) + 1) * DEFAULT_TAB_PT
                elif child.tag in (f'{_W}br', f'{_W}cr') and child.get(f'{_W}type') in (None, 'textWrapping'):
                    place_word()
                    lines, x = lines + 1, 0.0
                elif child.tag == f'{_W}drawing':
                    extent = child.find(f'{_WP}inline/{_WP}extent')
                    if extent is not None and str(extent.get('cy', '')).isdigit():
                        place_word()
                        picture_h = max(picture_h, int(extent.get('cy')) / EMU_PER_PT)
                        word = int(extent.get('cx', 0) or 0) / EMU_PER_PT
                        place_word()
        place_word()

        base = text_line or mark_line
        line_value, rule = spacing('line'), spacing('lineRule') or 'auto'
        if rule == 'exact':
            line = _twips(line_value, base)
        elif rule == 'atLeast':
            line = max(base, _twips(line_value))
        else:
            line = base * (int(line_value) / 240 if line_value and line_value.isdigit() else 1.0)
        height = lines * line
        if picture_h:
            height = max(height, picture_h + (lines - 1) * line)

        border_key, edges = '', {'top': 0.0, 'bottom': 0.0}
        borders = self._first(p_chain, 'pBdr')
        if borders is not None:
            border_key = etree.tostring(borders).decode()
            for side in edges:
                edge = borders.find(f'{_W}{side}')
                if edge is not None and edge.get(f'{_W}val') not in (None, 'nil', 'none'):
                    sz, space = edge.get(f'{_W}sz', '4'), edge.get(f'{_W}space', '0')
                    edges[side] = (int(sz) / 8 if sz.isdigit() else 0.5) + (int(space) if space.isdigit() else 0)
        keep = self._first(p_chain, 'keepNext')
        return _Box(before=_twips(spacing('before')), height=height + edges['top'] + edges['bottom'],
                    after=_twips(spacing('after')), line=line or base, keep_next=keep is not None and _on(keep),
                    border_key=border_key, border_top=edges['top'], border_bottom=edges['bottom'])

    def row_layout(self, table, row, width: float) -> Tuple[float, Dict]:
        """(height, {cell element: (text width, [top of each paragraph])}) of a table row."""
        from docx.table import _Cell
        tbl = table._tbl
        grid = [_twips(col.get(f'{_W}w')) for col in tbl.findall(f'{_W}tblGrid/{_W}gridCol')]
        def cell_margins(cell_mar, margins: Dict[str, float]) -> Dict[str, float]:
            margins = dict(margins)
            for side, names in (('left', ('left', 'start')), ('right', ('right', 'end')),
                                ('top', ('top',)), ('bottom', ('bottom',))):
                for name in names:
                    element = cell_mar.find(f'{_W}{name}') if cell_mar is not None else None
                    if element is not None and element.get(f'{_W}type', 'dxa') == 'dxa':
                        margins[side] = _twips(element.get(f'{_W}w'), margins[side])
                        break
            return margins

        table_margins = cell_margins(tbl.find(f'{_W}tblPr/{_W}tblCellMar'), {
            'left': DEFAULT_CELL_MARGIN_PT, 'right': DEFAULT_CELL_MARGIN_PT, 'top': 0.0, 'bottom': 0.0})

        tr = row._tr
        tr_pr = tr.trPr
        column = 0
        if tr_pr is not None and tr_pr.find(f'{_W}gridBefore') is not None:
            column = int(tr_pr.find(f'{_W}gridBefore').get(f'{_W}val', 0) or 0)
        tcs = tr.tc_lst
        height, cells = 0.0, {}
        for tc in tcs:
            span = tc.grid_span
            cell_w = sum(grid[column:column + span]) if column + span <= len(grid) else 0.0
            column += span
            if cell_w <= 0:
                cell_w = width / max(len(tcs), 1)
            margins = cell_margins(tc.find(f'{_W}tcPr/{_W}tcMar'), table_margins)
            avail = max(18.0, cell_w - margins['left'] - margins['right'])
            cell = _Cell(tc, table)
            tops, content = _stack([self.paragraph_box(p, avail) for p in cell.paragraphs])
            content += sum(self.row_layout(nested, nested_row, avail)[0]
                           for nested in cell.tables for nested_row in nested.rows)
            height = max(height, margins['top'] + content + margins['bottom'])
            cells[tc] = (avail, [margins['top'] + top for top in tops])

        tr_height = tr_pr.find(f'{_W}trHeight') if tr_pr is not None else None
        if tr_height is not None:
            fixed = _twips(tr_height.get(f'{_W}val'))
            height = fixed if tr_height.get(f'{_W}hRule') == 'exact' else max(height, fixed)
        return height, cells

    @staticmethod
    def is_header_row(row) -> bool:
        tr_pr = row._tr.trPr
        header = tr_pr.find(f'{_W}tblHeader') if tr_pr is not None else None
        return header is not None and _on(header)

    def story_layout(self, story, width: float) -> Tuple[float, Dict]:
        """(height, {paragraph element: top offset}) of a header / footer."""
        from docx.table import Table
        from docx.text.paragraph import Paragraph
        offset, previous, tops = 0.0, None, {}
        for block in story.iter_inner_content():
            if isinstance(block, Paragraph):
                box = self.paragraph_box(block, width)
                offset += _gap(previous, box)
                tops[block._p] = offset
                offset += box.height + box.after
                previous = box
            elif isinstance(block, Table):
                offset += sum(self.row_layout(block, row, width)[0] for row in block.rows)
                previous = None
        return offset, tops


def docx_layout(document) -> TemplateLayout:
    """
    Pages follow the page breaks in the file: explicit breaks and the breaks Word last rendered.
    Files without rendered breaks are paginated by their estimated text flow.
    """
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    section = document.sections[0]
    page_w = section.page_width.pt if section.page_width else 612.0
    page_h = section.page_height.pt if section.page_height else 792.0
    m_left = section.left_margin.pt if section.left_margin is not None else 72.0
    m_right = section.right_margin.pt if section.right_margin is not None else 72.0
    m_top = section.top_margin.pt if section.top_margin is not None else 72.0
    m_bottom = section.bottom_margin.pt if section.bottom_margin is not None else 72.0
    header_distance = section.header_distance.pt if section.header_distance is not None else 36.0
    footer_distance = section.footer_distance.pt if section.footer_distance is not None else 36.0
    text_w = max(72.0, page_w - m_left - m_right)
    flow = _DocxFlow(document)

    def story(kind: str, first: bool):
        """Header / footer of the first or the following pages; None when the section has none."""
        try:
            first = first and section.different_first_page_header_footer
            part = getattr(section, f'first_page_{kind}' if first else kind)
            return None if part.is_linked_to_previous else part  # linked = no header of its own
        except Exception:
            logger.debug("DOCX %s could not be read", kind, exc_info=True)
            return None

    def text_area(first: bool) -> Tuple[float, float]:
        """Top and bottom of the body text; a tall header / footer pushes it in from the margins."""
        top, bottom = m_top, page_h - m_bottom
        header, footer = story('header', first), story('footer', first)
        try:
            if header is not None:
                top = max(top, header_distance + flow.story_layout(header, text_w)[0])
            if footer is not None:
                bottom = min(bottom, page_h - footer_distance - flow.story_layout(footer, text_w)[0])
        except Exception:
            logger.debug("DOCX header / footer height could not be estimated", exc_info=True)
        return top, max(bottom, top + 72.0)

    areas = (text_area(True), text_area(False))

    def area(page_no: int) -> Tuple[float, float]:
        return areas[0] if page_no == 1 else areas[1]

    # Word's own page breaks are exact; without them the overflow of the text area starts a page
    paginate = document.element.body.find(f'.//{_W}lastRenderedPageBreak') is None

    layout = TemplateLayout(page_count=1, page_width=page_w, page_height=page_h)
    # (page, text, align, pictures[(mime, data, w, h, x, y, follows_text)], y, bottom)
    entries: List[Tuple[int, str, str, list, float, float]] = []
    page = 1
    cursor = area(1)[0]
    previous: Optional[_Box] = None  # the body paragraph just above the cursor
    content_since_break = False

    def next_page():
        nonlocal page, cursor, previous, content_since_break
        page += 1
        cursor = area(page)[0]
        previous = None
        content_since_break = False

    def picture_entries(element, part, align, paragraph_y: float) -> list:
        found = []
        for drawing in element.iter(f'{_W}drawing'):
            picture = _docx_picture(drawing, part)
            if not picture:
                continue
            mime, data, w_pt, h_pt = picture
            anchor = drawing.find(f'{_WP}anchor')
            x_pt = y_pt = None
            follows_text = True
            if anchor is not None:
                x_pt = _anchor_offset(anchor, 'H', w_pt, page_w, m_left, m_right)
                y_pt = _anchor_offset(anchor, 'V', h_pt, page_h, m_top, m_bottom, paragraph_y)
                position = anchor.find(f'{_WP}positionV')
                follows_text = y_pt is None or (position is not None and
                                                position.get('relativeFrom') in ('paragraph', 'line'))
            if x_pt is None:
                if align == 'center':
                    x_pt = (page_w - w_pt) / 2
                elif align == 'right':
                    x_pt = page_w - m_right - w_pt
                else:
                    x_pt = m_left
            found.append((mime, data, w_pt, h_pt, x_pt, paragraph_y if y_pt is None else y_pt, follows_text))
        return found

    def visit_paragraph(paragraph: Paragraph, width: float, cell_top: Optional[float] = None):
        """Add the paragraph; inside a table row it is placed at `cell_top`."""
        nonlocal page, cursor, previous, content_since_break
        in_table = cell_top is not None
        p = paragraph._p
        p_pr = p.pPr
        if p_pr is not None and p_pr.find(f'{_W}pageBreakBefore') is not None and content_since_break:
            next_page()
        box = flow.paragraph_box(paragraph, width)
        split_at = None
        if in_table:
            y = cell_top
        else:
            gap = _gap(previous, box)
            bottom = area(page)[1]
            keep = box.line if box.keep_next else 0.0
            if paginate and content_since_break and cursor + gap + box.height + keep > bottom:
                # Long paragraphs break across the pages (two lines at least on each side)
                lines = round(box.height / box.line) if box.line else 1
                fits = int((bottom - cursor - gap) // box.line) if box.line else 0
                if not box.keep_next and lines >= 4 and fits >= 2 and lines - fits >= 2:
                    split_at = cursor + gap + fits * box.line
                else:
                    next_page()
                    gap = 0.0  # space before is dropped at the top of a page
            y = cursor + gap
        start_page = page
        text_before_break = ''
        broke_inside = False
        for node in p.iter(f'{_W}t', f'{_W}br', f'{_W}lastRenderedPageBreak'):
            if node.tag == f'{_W}t':
                text_before_break += node.text or ''
                content_since_break = content_since_break or bool((node.text or '').strip())
            elif (node.tag == f'{_W}br' and node.get(f'{_W}type') == 'page') or \
                    node.tag == f'{_W}lastRenderedPageBreak':
                if content_since_break:
                    page += 1
                    broke_inside = bool(text_before_break.strip())
                    content_since_break = False
        text = paragraph.text.strip()
        align = _docx_align(paragraph)
        # Text before an in-paragraph break stays on the previous page
        entry_page = page - 1 if broke_inside and page > 1 else page
        if page != start_page and not broke_inside:
            y = area(page)[0]  # the paragraph starts on the new page
        entries.append((entry_page, text, align, picture_entries(p, paragraph.part, align, y), y, y + box.height))
        if text and not in_table:
            prefix = f'[{align}] ' if align in ('center', 'right') else ''
            layout.llm_parts.append((entry_page, prefix + text))
        if text or entries[-1][3]:
            content_since_break = True
        if not in_table:
            if page != start_page:
                cursor = area(page)[0] + (box.line if broke_inside else box.height) + box.after
            elif split_at is not None:
                next_page()
                cursor += y + box.height - split_at + box.after
                content_since_break = True
            else:
                cursor = y + box.height + box.after
            previous = box
        if p_pr is not None and p_pr.find(f'{_W}sectPr') is not None:
            next_page()  # section break (next page)

    for block in document.iter_inner_content():
        if isinstance(block, Paragraph):
            visit_paragraph(block, text_w)
        elif isinstance(block, Table):
            previous = None
            header_rows, leading = 0.0, True
            for row in block.rows:
                row_h, cell_boxes = flow.row_layout(block, row, text_w)
                if paginate and content_since_break and cursor + row_h > area(page)[1]:
                    next_page()
                    cursor += header_rows  # Word repeats the header rows on the new page
                leading = leading and flow.is_header_row(row)
                if leading:
                    header_rows += row_h
                row_page, row_top = page, cursor
                for cell in row.cells:
                    width, tops = cell_boxes.get(cell._tc) or \
                        (text_w, _stack([flow.paragraph_box(p, text_w) for p in cell.paragraphs])[0])
                    for paragraph, top in zip(cell.paragraphs, tops):
                        visit_paragraph(paragraph, width, cell_top=row_top + top)
                cursor = row_top + row_h if page == row_page else area(page)[0] + row_h
                cells = [cell.text.strip().replace('\n', ' / ') for cell in row.cells]
                if any(cells):
                    layout.llm_parts.append((row_page, ' | '.join(cells)))

    layout.page_count = max([page] + [e[0] for e in entries]) if entries else 1
    # Drop a trailing page that only exists because of a final break
    while layout.page_count > 1 and not any(e[0] == layout.page_count and (e[1] or e[3]) for e in entries):
        layout.page_count -= 1

    # Estimated heights can overrun a page Word ended itself; squeeze such pages into the text area
    squeeze: Dict[int, Tuple[float, float]] = {}
    for page_no in range(1, layout.page_count + 1):
        top, bottom = area(page_no)
        end = max((e[5] for e in entries if e[0] == page_no), default=top)
        if end > bottom and end > top:
            squeeze[page_no] = (top, (bottom - top) / (end - top))

    def fit(page_no: int, y_pt: float) -> float:
        if page_no not in squeeze or y_pt <= squeeze[page_no][0]:
            return y_pt
        top, factor = squeeze[page_no]
        return top + (y_pt - top) * factor

    for entry_page, text, align, pictures, y_pt, _ in entries:
        if text:
            layout.lines.append(LayoutLine(page=entry_page, y=_clamp(fit(entry_page, y_pt) / page_h), text=text,
                                           align=align, norm=_norm(text)))
        for mime, data, w_pt, h_pt, x_pt, pic_y, follows_text in pictures:
            if follows_text:
                pic_y = fit(entry_page, y_pt) + (pic_y - y_pt)
            _add_image(layout, LayoutImage(
                page=entry_page, x=_clamp(x_pt / page_w), y=_clamp(pic_y / page_h),
                w=_clamp(w_pt / page_w), h=_clamp(h_pt / page_h), mime=mime, data=data,
            ))
    layout.measured_pages = set(range(1, layout.page_count + 1))

    # Header pictures (logos) repeat on every page
    try:
        for page_no in range(1, layout.page_count + 1):
            header = story('header', page_no == 1)
            if header is None:
                continue
            tops = flow.story_layout(header, text_w)[1]
            for paragraph in header.paragraphs:
                align = _docx_align(paragraph)
                paragraph_y = header_distance + tops.get(paragraph._p, 0.0)
                for mime, data, w_pt, h_pt, x_pt, y_pt, _ in picture_entries(paragraph._p, header.part, align,
                                                                             paragraph_y):
                    _add_image(layout, LayoutImage(
                        page=page_no, x=_clamp(x_pt / page_w), y=_clamp(y_pt / page_h),
                        w=_clamp(w_pt / page_w), h=_clamp(h_pt / page_h), mime=mime, data=data,
                    ))
    except Exception:
        logger.debug("Header picture extraction failed", exc_info=True)
    return layout


def docx_text_for_llm(layout: TemplateLayout) -> str:
    """DOCX text with page markers and the alignment of centred / right-aligned lines; table rows as "a | b"."""
    parts, current = [], None
    for page, text in layout.llm_parts:
        if page != current:
            current = page
            parts.append(f'--- Page {current} ---')
        parts.append(text)
    return '\n'.join(parts)


# ------------------------------------------------------------- placement

def _match_quality(heading: str, target: str, line: LayoutLine) -> int:
    """3: the line is the heading (or "Heading: value"); 2: starts with it; 1: contains it; 0: no match."""
    if line.norm == target:
        return 3
    if line.norm.startswith(target):
        label = re.escape(' '.join(heading.lower().split()))
        if re.match(rf'^\W*{label}\s*[:\-–]', ' '.join(line.text.lower().split())):
            return 3
        return 2
    return 1 if target in line.norm else 0


def place_fields(fields: List[Dict], layout: TemplateLayout) -> None:
    """
    Set `page`, `pos_y` (and `align` for measured text) on each field, in place.
    Fields are in document order, so matching continues forward from the previous match.
    """
    cursor = 0  # index into layout.lines
    last_page, last_y = 1, 0.0
    lines = layout.lines
    for index, item in enumerate(fields):
        # Columns of a table share the placement of the table's first column
        if item.get('table_key') is not None and index > 0 and fields[index - 1].get('table_key') == item['table_key']:
            item['page'], item['pos_y'] = fields[index - 1]['page'], fields[index - 1]['pos_y']
            continue

        target = _norm(item['heading'])
        match = None
        if target:
            # Best kind of match wins over an earlier weaker one, e.g. the title line "Report" over a
            # running header "REPORT - Client X" that merely starts with the same words
            best = 0
            for j in range(cursor, len(lines)):
                quality = _match_quality(item['heading'], target, lines[j])
                if quality > best:
                    best, match = quality, j
                    if quality == 3:
                        break
        if match is not None:
            line = lines[match]
            cursor = match + 1
            item['page'], item['pos_y'] = line.page, line.y
            if item.get('table_key') is None and line.align:
                item['align'] = line.align if line.align in ('center', 'right') else None
        else:
            # Not measurable (scanned page): the LLM's estimate, kept in document order
            page = item.get('llm_page')
            page = page if isinstance(page, int) and last_page <= page <= layout.page_count else last_page
            y = item.get('llm_y')
            y = y if isinstance(y, (int, float)) and 0 <= y <= 1 else None
            if y is None or (page == last_page and y < last_y):
                y = last_y
            item['page'], item['pos_y'] = page, y
        last_page, last_y = item['page'], item['pos_y']
